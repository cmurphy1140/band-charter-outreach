"""Team-facing outputs from data/final/prospects.csv.

  python scripts/export.py            # prospects.xlsx + SUMMARY.md
  python scripts/export.py --gsheet SERVICE_ACCOUNT.json [--gsheet-title "Parade Prospects Tier A"]

Only the scrapers and public sites feed these files; no email or inbox data.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from scrapers.common import COLUMNS, FINAL_DIR, is_east_coast  # noqa: E402

PROSPECTS = FINAL_DIR / "prospects.csv"
XLSX = FINAL_DIR / "prospects.xlsx"
SUMMARY = FINAL_DIR / "SUMMARY.md"
EAST_COAST_CSV = FINAL_DIR / "east_coast.csv"

# Column order for the workbook: the "Parade" column is the parades list, one per line.
SHEET_COLUMNS = [
    ("Tier", "tier"), ("Score", "score"), ("School", "school"), ("Band", "band_name"),
    ("City", "city"), ("State", "state"), ("Level", "level"), ("Parade", "parades"), ("Parades marched", "parades_marched"),
    ("Last appearance", "last_appearance"), ("BOA finalist years", "boa_finalist_years"),
    ("District", "district"), ("Enrollment", "enrollment"), ("School site", "school_url"),
    ("Band site", "band_url"), ("Director", "director_name"), ("Director email", "director_email"),
    ("Director phone", "director_phone"), ("Boosters", "booster_org"), ("Sources", "source_urls"),
    ("Notes", "notes"),
]
WIDTHS = {"School": 34, "Band": 28, "Parade": 36, "District": 26, "School site": 30, "Band site": 34,
          "Director": 20, "Director email": 28, "Boosters": 26, "Sources": 40, "Notes": 40}


def load_rows() -> list[dict]:
    with PROSPECTS.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    rows.sort(key=lambda r: (-float(r["score"] or 0), r["school"]))
    return rows


def _write_sheet(ws, rows: list[dict]) -> None:
    ws.append([h for h, _ in SHEET_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        vals = []
        for _, key in SHEET_COLUMNS:
            v = r.get(key, "")
            if key == "parades":
                v = "\n".join(p for p in v.split("; ") if p)
            elif key == "source_urls":
                v = "\n".join(p for p in v.split("; ") if p)
            elif key in ("score", "parades_marched", "last_appearance", "enrollment") and v not in ("", None):
                try:
                    v = float(v) if key == "score" else int(float(v))
                except ValueError:
                    pass
            vals.append(v)
        ws.append(vals)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SHEET_COLUMNS))}{max(ws.max_row, 1)}"
    for i, (h, key) in enumerate(SHEET_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 14)
        if key in ("parades", "source_urls", "notes"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tier A"
    _write_sheet(ws, [r for r in rows if r["tier"] == "A"])
    _write_sheet(wb.create_sheet("Tier B"), [r for r in rows if r["tier"] == "B"])
    east = [r for r in rows if is_east_coast(r)]
    _write_sheet(wb.create_sheet("East Coast"), east)
    _write_sheet(wb.create_sheet("All"), rows)
    wb.save(XLSX)
    with EAST_COAST_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(east)


def write_summary(rows: list[dict]) -> None:
    by_state = Counter(r["state"] or "unknown" for r in rows)
    by_parade = Counter()
    for r in rows:
        for tag in r["parades"].split("; "):
            if tag:
                by_parade[tag.rsplit(" ", 1)[0]] += 1
    tiers = Counter(r["tier"] for r in rows)
    with_email = sum(1 for r in rows if r["director_email"])
    no_contact = [r for r in rows if not r["director_email"] and not r["director_phone"]]
    lines = [
        "# Parade prospects summary", "",
        f"_Generated {dt.date.today().isoformat()} from `data/final/prospects.csv`._", "",
        f"- **{len(rows)} schools** (Tier A {tiers['A']}, Tier B {tiers['B']}, Tier C {tiers['C']})",
        f"- {with_email} with a published director email; {len(no_contact)} with no contact found",
        "", "## By state", "", "| State | Schools |", "|---|---|",
    ]
    lines += [f"| {s} | {n} |" for s, n in sorted(by_state.items(), key=lambda x: (-x[1], x[0]))]
    lines += ["", "## By parade", "", "| Parade | Schools |", "|---|---|"]
    lines += [f"| {p} | {n} |" for p, n in by_parade.most_common()]
    lines += ["", "## Top 25 schools", "", "| # | School | State | Score | Parade history |", "|---|---|---|---|---|"]
    for i, r in enumerate(rows[:25], 1):
        lines.append(f"| {i} | {r['school']} | {r['state'] or '?'} | {r['score']} | {r['parades']} |")
    lines += ["", "## Schools with no contact found (manual lookup)", "",
              "Director contacts are recorded only when published on the school, district, or",
              "booster site. These schools need a manual lookup:", "",
              "| School | State | Tier | School site | Band site |", "|---|---|---|---|---|"]
    for r in no_contact:
        lines.append(f"| {r['school']} | {r['state'] or '?'} | {r['tier']} | {r['school_url']} | {r['band_url']} |")
    SUMMARY.write_text("\n".join(lines) + "\n", encoding="utf-8")


def push_gsheet(rows: list[dict], service_account_json: str, title: str) -> str:
    try:
        import gspread  # noqa: F401
    except ImportError:
        sys.exit("gspread is not installed. It is not in requirements.txt by design; "
                 "install it in the venv (pip install gspread) once a service account is provided.")
    import gspread
    gc = gspread.service_account(filename=service_account_json)
    sh = gc.create(title)
    ws = sh.sheet1
    ws.update_title("Tier A")
    tier_a = [r for r in rows if r["tier"] == "A"]
    values = [[h for h, _ in SHEET_COLUMNS]]
    for r in tier_a:
        values.append([("\n".join(r[k].split("; ")) if k in ("parades", "source_urls") else r.get(k, ""))
                       for _, k in SHEET_COLUMNS])
    ws.update(values)
    ws.freeze(rows=1)
    return sh.url


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--gsheet", metavar="SERVICE_ACCOUNT_JSON",
                    help="push the Tier A sheet to a new Google Sheet with this service account")
    ap.add_argument("--gsheet-title", default="Parade Prospects - Tier A")
    a = ap.parse_args(argv)
    rows = load_rows()
    write_xlsx(rows)
    write_summary(rows)
    print(f"wrote {XLSX.relative_to(FINAL_DIR.parent.parent)} ({len(rows)} rows) and "
          f"{SUMMARY.relative_to(FINAL_DIR.parent.parent)}")
    if a.gsheet:
        print("google sheet:", push_gsheet(rows, a.gsheet, a.gsheet_title))
    return 0


if __name__ == "__main__":
    sys.exit(main())
