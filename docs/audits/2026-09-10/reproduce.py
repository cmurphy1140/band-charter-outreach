import csv,json,sys,random,os,re
from pathlib import Path
from unittest.mock import patch
from tempfile import TemporaryDirectory
from collections import Counter
from bs4 import BeautifulSoup
import requests
ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT))
from scripts import run_all,enrich,score,export
from scrapers import common,serpapi,boa,heb
from scrapers.normalize import normalize_school
from openpyxl import Workbook,load_workbook
rows=list(csv.DictReader((ROOT/'data/final/prospects.csv').open()))
result={}
result['baseline_commit']=__import__('subprocess').check_output(['git','rev-parse','HEAD'],cwd=ROOT).decode().strip()
result['counts']={f:sum(bool(r.get(f)) for r in rows) for f in common.COLUMNS}
result['tiers']=dict(Counter(r['tier'] for r in rows))
result['future_rows']=sum(int(r['last_appearance'] or 0)>2026 for r in rows)
result['boa_only_rows']=sum(all(p.startswith('BOA') for p in r['parades'].split('; ')) for r in rows)
result['scoring_mismatches']=sum(float(r['score'])!=score.score_row(r)[0] for r in rows)
result['sample']=random.Random(20260910).sample(rows,15)
result['search_placements']=[{k:r[k] for k in ('school','city','state','school_url','notes')} for r in rows if 'state from Google' in r['notes'] or 'unique national match' in r['notes']]
result['organic_picks']=[{k:r[k] for k in ('school','city','state','school_url','notes')} for r in rows if 'website via Google result' in r['notes']]
wb=load_workbook(ROOT/'data/final/prospects.xlsx',read_only=True)
result['xlsx_rows']={ws.title:ws.max_row-1 for ws in wb}
result['xlsx_formulas']=sum(c.data_type=='f' for ws in wb for row in ws for c in row)
result['formula_inputs']=[(r['school'],k) for r in rows for k,v in r.items() if v.startswith(('=','+','@','\t','\r'))]
html='<div>Jane Smith, Theatre Director jane@district.org John Brown, Band Director john@district.org</div>'
soup=BeautifulSoup(html,'lxml'); found=dict(director_email='',director_name='',director_phone='',booster_org='',social=[],contact_source='')
enrich._extract_contacts(soup,soup.get_text(' ',strip=True),'https://district.org/arts','district.org',found)
result['adjacent_contact_probe']=found
result['domain_collision']=[enrich._domain(u) for u in ['https://a.k12.tx.us','https://b.k12.tx.us','https://school1.edublogs.org','https://school2.edublogs.org']]

from unittest.mock import Mock
import pandas as pd
nces_probe = pd.DataFrame([dict(ncessch=str(i), sch_name='Lincoln High School', key='lincoln', state='TX', city=city, city_key=city.lower(), level='High') for i, city in enumerate(['Austin', 'Dallas'])])
matched, ratio, others = enrich.match_nces(dict(school='Lincoln High School',state='TX',city='Unmatched city'),nces_probe)
result['unmatched_city_tie']={'selected_city':matched['city'],'ratio':ratio,'other_candidates':others}
with patch.dict(common._robots,{},clear=True),patch.object(common,'_throttle'),patch.object(common._session,'get',return_value=Mock(status_code=503)):
    result['robots_503_allowed']=common._robots_allowed('https://example.org/path')
new_rows, _ = run_all.merge(run_all.load_interim())
new_rows = run_all.carry_over(new_rows, rows)
row_key = lambda r: (normalize_school(r['school']),r['state'])
original = {row_key(r): {k:str(v) for k,v in r.items()} for r in rows}
rebuilt = {row_key(r): {k:str(v) for k,v in r.items()} for r in new_rows}
result['current_roundtrip_identical']=original==rebuilt

with TemporaryDirectory() as td:
    p=Path(td)
    with patch.object(common,'INTERIM_DIR',p),patch.object(run_all,'INTERIM_DIR',p):
        common.write_interim('probe',[common.Row('Lincoln High School',state='TX',event='Rose',year=2026,source_url='https://example.org/tx'),common.Row('Lincoln High School',state='CA',event='Rose',year=2026,source_url='https://example.org/ca')])
        incoming=[common.Row('Lincoln High School',state='TX',event='Rose',year=2026,source_url='https://example.org/tx'),common.Row('Lincoln High School',state='CA',event='Rose',year=2026,source_url='https://example.org/ca')]
        run_all.write_interim_scoped('probe',incoming,{2026})
        result['same_name_states_kept']=list(csv.DictReader((p/'probe.csv').open()))
        common.write_interim('boa',[common.Row('Existing High School',state='TX',event='BOA Grand National Finalist',year=2026,source_url='https://example.org')])
        with patch.object(boa,'fetch',side_effect=common.BlockedSource('simulated source outage')):
            returned=boa.scrape(years={2026,2027})
        run_all.write_interim_scoped('boa',returned,{2026,2027})
        result['blocked_boa_rows_remaining']=len(list(csv.DictReader((p/'boa.csv').open())))
    with patch.object(common,'RAW_DIR',p):
        url='https://example.org/lineup'; cp=common.cache_path(url);cp.parent.mkdir(parents=True);cp.write_text('old lineup')
        with patch.object(common._session,'get') as get:
            result['cache_return']=common.fetch(url);result['cache_network_calls']=get.call_count
    fake_key='AUDIT_SYNTHETIC_KEY_NOT_REAL'
    with patch.dict(os.environ,{'SERPAPI_KEY':fake_key}),patch.object(serpapi,'_throttle'),patch.object(serpapi.requests,'get',side_effect=requests.ConnectionError('failed https://serpapi.com/search.json?api_key='+fake_key)):
        try: serpapi.search('google','audit unique synthetic')
        except common.BlockedSource as exc: result['exception_contains_synthetic_key']=fake_key in str(exc)
    out=Workbook(); export._write_sheet(out.active,[{'school':'=1+1'}]); result['export_formula_type']=out.active['C2'].data_type
    with patch.object(run_all,'CHANGELOG',p/'CHANGELOG.md'):
        from contextlib import redirect_stdout
        import io
        buf=io.StringIO()
        with redirect_stdout(buf):run_all.diff_and_log([rows[0]],[])
        result['removed_row_diff']=buf.getvalue().strip()
# Key scan reports only counts/paths, never bytes or matched strings.
tracked=__import__('subprocess').check_output(['git','ls-files','-z'],cwd=ROOT).decode().split('\0')
key=os.environ.get('SERPAPI_KEY','');result['live_key_available']=bool(key)
result['exact_live_key_files']=[f for f in tracked if f and key and key.encode() in (ROOT/f).read_bytes()]
patterns=[rb'api_key["\s]*[:=]["\s]*[A-Za-z0-9_-]{20,}',rb'gh[pousr]_[A-Za-z0-9]{30,}',rb'-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----']
result['credential_pattern_files']=[f for f in tracked if f and any(re.search(p,(ROOT/f).read_bytes()) for p in patterns)]
Path(__file__).with_name('evidence.json').write_text(json.dumps(result,indent=2))
print(json.dumps({k:v for k,v in result.items() if k not in ('sample','search_placements','organic_picks')},indent=2))
